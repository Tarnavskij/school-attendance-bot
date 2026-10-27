# helpers/excel_export.py
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


def generate_absences_excel(data: dict, title_info: str = "Все классы") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Пропуски"

    # Заголовок отчета
    ws.cell(row=1, column=1, value=f"Отчет по пропускам: {title_info}").font = Font(bold=True, size=12)

    headers = ["Класс", "ФИО", "Всего", "По болезни", "На оздоровлении", "По заявлению", "Без уважительной"]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    row_num = 4
    for class_name, students in data.items():
        for s in students:
            ws.cell(row=row_num, column=1, value=class_name)
            ws.cell(row=row_num, column=2, value=s["name"])
            ws.cell(row=row_num, column=3, value=s["total"])
            ws.cell(row=row_num, column=4, value=s["sick"])
            ws.cell(row=row_num, column=5, value=s["recuperation"])
            ws.cell(row=row_num, column=6, value=s["statement"])
            ws.cell(row=row_num, column=7, value=s["unexcused"])
            row_num += 1

    for col in ws.columns:
        width = max((len(str(cell.value or "")) for cell in col), default=10) + 2
        ws.column_dimensions[col[0].column_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()